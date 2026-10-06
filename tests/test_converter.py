"""Tests for the LibreOffice conversion service."""

import shutil
from pathlib import Path

import pytest
from docx import Document

from app.converter import ConversionError, convert_to_pdf

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session", autouse=True)
def ensure_fixtures() -> None:
    """Generate a small .docx fixture once per test session."""
    FIXTURES.mkdir(exist_ok=True)
    docx_path = FIXTURES / "sample.docx"
    if not docx_path.exists():
        doc = Document()
        doc.add_heading("Word to PDF", level=1)
        doc.add_paragraph("This is a sample document for testing.")
        doc.add_paragraph("It should convert cleanly to PDF.")
        doc.save(docx_path)


@pytest.fixture
def sample_docx() -> Path:
    return FIXTURES / "sample.docx"


@pytest.mark.skipif(
    shutil.which("soffice") is None and shutil.which("libreoffice") is None,
    reason="LibreOffice not installed",
)
@pytest.mark.asyncio
async def test_convert_docx_to_pdf(sample_docx: Path, tmp_path: Path) -> None:
    """A .docx should convert to a valid PDF."""
    pdf = await convert_to_pdf(sample_docx, output_dir=tmp_path)

    assert pdf.exists()
    assert pdf.suffix == ".pdf"
    assert pdf.stat().st_size > 0

    # PDFs start with "%PDF-" magic bytes.
    with pdf.open("rb") as f:
        assert f.read(5) == b"%PDF-"


@pytest.mark.asyncio
async def test_convert_missing_file_raises(tmp_path: Path) -> None:
    """A missing input file should raise ConversionError."""
    with pytest.raises(ConversionError):
        await convert_to_pdf(tmp_path / "does-not-exist.docx", output_dir=tmp_path)
